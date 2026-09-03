import os
import re
import json
import uuid
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.utils.logger import logger

EXPORTS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "exports"))
os.makedirs(EXPORTS_DIR, exist_ok=True)

def _clean_cell_value(val: Any) -> Any:
    if val is None:
        return ""
    if isinstance(val, dict):
        if "$oid" in val:
            return val["$oid"]
        if "$date" in val:
            return val["$date"]
        return json.dumps(val)
    if isinstance(val, (list, tuple)):
        return json.dumps(val)
    return val

def _format_header_label(key: str) -> str:
    # Convert snake_case or camelCase to Title Case
    # e.g., 'vehicle_no' -> 'Vehicle No', 'runDate' -> 'Run Date'
    s1 = re.sub('(.)([A-Z][a-z]+)', r'\1 \2', key)
    s2 = re.sub('([a-z0-9])([A-Z])', r'\1 \2', s1)
    s3 = s2.replace('_', ' ').strip()
    return ' '.join(word.capitalize() for word in s3.split())

def generate_trips_excel(records: List[Dict[str, Any]], filename_prefix: str = "trips_export") -> str:
    """
    Generates a professionally styled Excel file (.xlsx) from a list of records.
    Returns the generated filename.
    """
    if not records or not isinstance(records, list):
        raise ValueError("No records available to export to Excel.")

    # Determine all unique headers across all records while preserving typical key order
    all_keys = []
    seen = set()
    for row in records:
        if isinstance(row, dict):
            for k in row.keys():
                if k not in seen:
                    seen.add(k)
                    all_keys.append(k)

    if not all_keys:
        all_keys = ["id", "data"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trips Report"

    # Header styling
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    # Write headers
    for col_num, key in enumerate(all_keys, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = _format_header_label(key)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write data rows
    for row_num, record in enumerate(records, 2):
        is_even = (row_num % 2 == 0)
        for col_num, key in enumerate(all_keys, 1):
            cell = ws.cell(row=row_num, column=col_num)
            val = record.get(key, "") if isinstance(record, dict) else ""
            cell.value = _clean_cell_value(val)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if not is_even:
                cell.fill = zebra_fill

    # Enable auto filter
    ws.auto_filter.ref = ws.dimensions

    # Adjust column widths nicely
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Add buffer and set max width
        adjusted_width = min(max(max_len + 4, 12), 45)
        ws.column_dimensions[col_letter].width = adjusted_width

    # Generate filename
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    unique_suffix = uuid.uuid4().hex[:6]
    filename = f"{filename_prefix}_{timestamp}_{unique_suffix}.xlsx"
    filepath = os.path.join(EXPORTS_DIR, filename)

    wb.save(filepath)
    logger.info(f"Generated Excel file: {filepath} ({len(records)} rows)")
    return filename
